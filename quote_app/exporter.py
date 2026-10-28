"""Quote workbook writer: drawing columns verbatim + RFQ template columns."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

TEMPLATE_HEADERS = ["单重", "总重", "单价", "总价"]
TEMPLATE_KEYS = ["unit_weight", "total_weight", "unit_price", "total_price"]

FONT = Font(name="宋体", size=11)
THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def collect_source_headers(items):
    headers = []
    for item in items:
        for name in (item.get("source") or {}):
            if name and name not in headers:
                headers.append(name)
    return headers


def _as_number(value):
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, str):
        text = value.strip().replace(",", "")
        if not text:
            return None
        try:
            number = float(text)
        except ValueError:
            return value
        return int(number) if number.is_integer() else number
    return value


def _source_width(name):
    text = name.lower()
    if any(k in text for k in ("description", "desc", "名称", "描述")):
        return 52
    if any(k in text for k in ("remark", "备注", "说明")):
        return 40
    return 13


def export_quote(items, path, source_headers=None):
    if source_headers is None:
        source_headers = collect_source_headers(items)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    headers = list(source_headers) + TEMPLATE_HEADERS
    for col, name in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = FONT
        cell.alignment = CENTER
        cell.border = BORDER

    numeric = {"unit_weight", "total_weight", "unit_price", "total_price"}
    for offset, item in enumerate(items, 1):
        row = 1 + offset
        ws.row_dimensions[row].height = 69
        source = item.get("source") or {}
        for col, name in enumerate(source_headers, 1):
            cell = ws.cell(row=row, column=col, value=source.get(name, ""))
            cell.font = FONT
            cell.alignment = CENTER
            cell.border = BORDER
        for tcol, key in enumerate(TEMPLATE_KEYS, 1):
            col = len(source_headers) + tcol
            value = item.get(key, "")
            if key in numeric:
                value = _as_number(value)
            cell = ws.cell(row=row, column=col, value=value)
            cell.font = FONT
            cell.alignment = CENTER
            cell.border = BORDER

    for col, name in enumerate(source_headers, 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = _source_width(name)
    template_widths = [12, 12, 12, 13]
    for tcol, width in enumerate(template_widths, 1):
        col = len(source_headers) + tcol
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
